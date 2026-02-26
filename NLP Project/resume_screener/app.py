import streamlit as st
import pandas as pd
import plotly.graph_objects as go
import plotly.express as px
from sentence_transformers import SentenceTransformer, util
import PyPDF2
from docx import Document
import io
from datetime import datetime
import json
from fpdf import FPDF
import re
from sklearn.feature_extraction.text import TfidfVectorizer
import numpy as np

# Page configuration
st.set_page_config(
    page_title="AI Resume Screener",
    page_icon="📄",
    layout="wide",
    initial_sidebar_state="expanded"
)

# Custom CSS for attractive UI
st.markdown("""
    <style>
    .main {
        background: linear-gradient(90deg, #141e30 0%, #243b55 100%)
    .stApp {
        background: linear-gradient(90deg, #141e30 0%, #243b55 100%)
    }
    .upload-box {
        background: #476f95;
        padding: 40px;
        border-radius: 20px;
        box-shadow: 0 8px 32px 0 rgba(31, 38, 135, 0.37);
        backdrop-filter: blur(4px);
        border: 1px solid rgba(255, 255, 255, 0.18);
        margin: 20px 0;
    }
    .upload-box h3 {
        color: #476f95;
        margin-top: 0;
        margin-bottom: 20px;
    }
    .stTabs [data-baseweb="tab-list"] {
        gap: 8px;
        background-color: rgba(255, 255, 255, 0.1);
        padding: 10px;
        border-radius: 15px;
    }
    .stTabs [data-baseweb="tab"] {
        background-color: rgba(255, 255, 255, 0.2);
        border-radius: 10px;
        padding: 12px 24px;
        color: white;
        font-weight: 600;
        font-size: 16px;
        border: 2px solid transparent;
        transition: all 0.3s ease;
    }
    .stTabs [data-baseweb="tab"]:hover {
        background-color: rgba(255, 255, 255, 0.3);
        border-color: rgba(255, 255, 255, 0.5);
    }
    .stTabs [aria-selected="true"] {
        background: linear-gradient(135deg, #60a5fa 0%, #a78bfa 100%) !important;
        border-color: white !important;
    }
    .tab-content {
        background: rgba(255, 255, 255, 0.95);
        padding: 30px;
        border-radius: 15px;
        margin-top: 20px;
        box-shadow: 0 8px 32px 0 rgba(31, 38, 135, 0.37);
    }
    .metric-card {
        background: linear-gradient(135deg, #f093fb 0%, #f5576c 100%);
        padding: 20px;
        border-radius: 15px;
        color: white;
        text-align: center;
        box-shadow: 0 4px 15px 0 rgba(245, 87, 108, 0.3);
    }
    .result-card {
        background: rgba(68, 24, 47, 0.8);
        padding: 25px;
        border-radius: 15px;
        margin: 15px 0;
        box-shadow: 0 4px 15px 0 rgba(0, 0, 0, 0.1);
    }
    .header-title {
        background: linear-gradient(135deg, #60a5fa 0%, #a78bfa 50%, #ec4899 100%);
        -webkit-background-clip: text;
        -webkit-text-fill-color: rgba(255, 255, 255, 0.8);
        font-size: 2.5em;
        font-weight: bold;
        text-align: center;
        
        font-family: monospace;
        letter-spacing: .17em;
        margin: 0 auto;
        overflow: hidden;
        white-space: nowrap;
        border-right: .17em solid rgba(255,255,255,0.75);
        
        animation: typing 3.5s steps(30, end), blinking-cursor .7s step-end infinite;
        
    }
    @keyframes typing {
        from {
            width: 0
            }
        to {
            width: 100%
        }
    }
    @keyframes blinking-cursor {
        from, 
        to {
            border-color: transparent
        }
        50% {
            border-color: rgba(255,255,255,0.75)
        }
    }
    .subtitle {
        text-align: center;
        color: white;
        font-size: 1.3em;
        margin-bottom: 30px;
    }
    .stButton>button {
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        border: none;
        padding: 12px 30px;
        border-radius: 25px;
        font-weight: bold;
        font-size: 24px !important;
        box-shadow: 0 4px 15px 0 rgba(102, 126, 234, 0.4);
        transition: all 0.3s ease;
    }
    .stButton>button:hover {
        transform: translateY(-2px);
        box-shadow: 0 6px 20px 0 rgba(102, 126, 234, 0.6);
    }
    .skill-badge {
        display: inline-block !important;
        margin-right: 8px;
        margin-bottom: 5px;
        background: linear-gradient(135deg, #667eea 0%, #764ba2 100%);
        color: white;
        padding: 5px 15px;
        border-radius: 20px;
        # margin: 5px;
        font-size: 14px;
    }
    </style>
""", unsafe_allow_html=True)

# Initialize session state
if 'analysis_done' not in st.session_state:
    st.session_state.analysis_done = False
if 'results' not in st.session_state:
    st.session_state.results = None

# Load BERT model
@st.cache_resource
def load_model():
    return SentenceTransformer('all-MiniLM-L6-v2')

model = load_model()

# File extraction functions
def extract_text_from_pdf(file):
    pdf_reader = PyPDF2.PdfReader(file)
    text = ""
    for page in pdf_reader.pages:
        text += page.extract_text()
    return text

def extract_text_from_docx(file):
    doc = Document(file)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text

def extract_text_from_txt(file):
    return file.read().decode('utf-8')

def extract_text(file):
    if file.name.endswith('.pdf'):
        return extract_text_from_pdf(file)
    elif file.name.endswith('.docx'):
        return extract_text_from_docx(file)
    elif file.name.endswith('.txt'):
        return extract_text_from_txt(file)
    return ""

# Skill extraction
def extract_skills(text):
    common_skills = [
        'python', 'java', 'javascript', 'c++', 'react', 'angular', 'vue',
        'machine learning', 'deep learning', 'data science', 'sql', 'nosql',
        'aws', 'azure', 'gcp', 'docker', 'kubernetes', 'git', 'agile',
        'scrum', 'project management', 'communication', 'leadership',
        'tensorflow', 'pytorch', 'sklearn', 'pandas', 'numpy', 'nlp',
        'computer vision', 'api', 'rest', 'html', 'css', 'node.js',
        'mongodb', 'postgresql', 'mysql', 'redis', 'kafka', 'spark'
    ]
    
    text_lower = text.lower()
    found_skills = []
    for skill in common_skills:
        if skill in text_lower:
            found_skills.append(skill.title())
    return found_skills

# Extract candidate details
def extract_candidate_details(resume_text):
    details = {}
    
    # Extract email
    email_pattern = r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b'
    emails = re.findall(email_pattern, resume_text)
    details['email'] = emails[0] if emails else 'Not found'
    
    # Extract phone
    phone_pattern = r'[\+]?[(]?[0-9]{3}[)]?[-\s\.]?[0-9]{3}[-\s\.]?[0-9]{4,6}'
    phones = re.findall(phone_pattern, resume_text)
    details['phone'] = phones[0] if phones else 'Not found'
    
    # Extract name (first line usually)
    lines = resume_text.split('\n')
    details['name'] = lines[0].strip() if lines else 'Not found'
    
    # Extract education keywords
    education_keywords = ['bachelor', 'master', 'phd', 'degree', 'university', 'college', 'mba']
    education = []
    for line in lines:
        if any(keyword in line.lower() for keyword in education_keywords):
            education.append(line.strip())
    details['education'] = education[:3] if education else ['Not found']
    
    # Extract years of experience
    exp_pattern = r'(\d+)\+?\s*(?:years?|yrs?)'
    exp_matches = re.findall(exp_pattern, resume_text.lower())
    if exp_matches:
        details['experience'] = f"{max(map(int, exp_matches))}+ years"
    else:
        details['experience'] = 'Not specified'
    
    # Word count
    details['word_count'] = len(resume_text.split())
    
    # Extract LinkedIn/GitHub
    linkedin_pattern = r'linkedin\.com/in/[\w-]+'
    github_pattern = r'github\.com/[\w-]+'
    linkedin = re.findall(linkedin_pattern, resume_text.lower())
    github = re.findall(github_pattern, resume_text.lower())
    details['linkedin'] = linkedin[0] if linkedin else 'Not found'
    details['github'] = github[0] if github else 'Not found'
    
    return details

# Analysis function
def analyze_resume(resume_text, job_description):
    # Encode texts
    resume_embedding = model.encode(resume_text, convert_to_tensor=True)
    job_embedding = model.encode(job_description, convert_to_tensor=True)
    
    # Calculate similarity
    similarity = util.pytorch_cos_sim(resume_embedding, job_embedding).item()
    match_percentage = similarity * 100
    
    # Extract skills
    resume_skills = extract_skills(resume_text)
    job_skills = extract_skills(job_description)
    
    # Skill match analysis
    matched_skills = list(set(resume_skills) & set(job_skills))
    missing_skills = list(set(job_skills) - set(resume_skills))
    
    # Extract candidate details
    candidate_details = extract_candidate_details(resume_text)
    
    # TF-IDF analysis for keyword matching
    vectorizer = TfidfVectorizer(max_features=20, stop_words='english')
    try:
        tfidf_matrix = vectorizer.fit_transform([resume_text, job_description])
        feature_names = vectorizer.get_feature_names_out()
        keywords = list(feature_names[:10])
    except:
        keywords = []
    
    # Determine recommendation
    if match_percentage >= 75:
        recommendation = "Highly Recommended"
        recommendation_color = "#00C853"
    elif match_percentage >= 60:
        recommendation = "Recommended"
        recommendation_color = "#FFB300"
    elif match_percentage >= 45:
        recommendation = "Fair"
        recommendation_color = "#FF6F00"
    else:
        recommendation = "Not Recommended"
        recommendation_color = "#D32F2F"
    
    return {
        'match_percentage': match_percentage,
        'recommendation': recommendation,
        'recommendation_color': recommendation_color,
        'resume_skills': resume_skills,
        'job_skills': job_skills,
        'matched_skills': matched_skills,
        'missing_skills': missing_skills,
        'keywords': keywords,
        'resume_text': resume_text,
        'job_description': job_description,
        'candidate_details': candidate_details
    }

# Visualization functions
def create_gauge_chart(value, title):
    fig = go.Figure(go.Indicator(
        mode = "gauge+number+delta",
        value = value,
        domain = {'x': [0, 1], 'y': [0, 1]},
        title = {'text': title, 'font': {'size': 24, 'color': '#333'}},
        delta = {'reference': 60},
        gauge = {
            'axis': {'range': [None, 100], 'tickwidth': 1, 'tickcolor': "#999"},
            'bar': {'color': "#297f94"},
            'bgcolor': "white",
            'borderwidth': 2,
            'bordercolor': "#ccc",
            'steps': [
                {'range': [0, 45], 'color': '#f5f5dc'},
                {'range': [45, 60], 'color': '#addfad'},
                {'range': [60, 75], 'color': '#98ff98'},
                {'range': [75, 100], 'color': '#39ff14'}
            ],
            'threshold': {
                'line': {'color': "red", 'width': 4},
                'thickness': 0.70,
                'value': 80
            }
        }
    ))
    
    fig.update_layout(
        paper_bgcolor='rgba(0,0,0,0)',
        plot_bgcolor='rgba(0,0,0,0)',
        font={'color': "#333", 'family': "Arial"},
        height=350
    )
    
    return fig

def create_skill_comparison_chart(matched, missing):
    labels = ['Matched Skills', 'Missing Skills']
    values = [len(matched), len(missing)]
    colors = ["#5166c3", '#f5576c']
    
    fig = go.Figure(data=[go.Pie(
        labels=labels,
        values=values,
        hole=0.4,
        marker=dict(colors=colors, line=dict(color='white', width=2)),
        textinfo='label+percent',
        textfont=dict(size=16, color='white'),
        hovertemplate='<b>%{label}</b><br>Count: %{value}<br>Percentage: %{percent}<extra></extra>'
    )])
    
    fig.update_layout(
        title={
            'text': 'Skill Match Analysis',
            'x': 0.5,
            'xanchor': 'center',
            'font': {'size': 24, 'color': '#333'}
        },
        paper_bgcolor='rgba(0,0,0,0)',
        plot_bgcolor='rgba(0,0,0,0)',
        font={'family': "Arial"},
        height=400,
        showlegend=True,
        legend=dict(
            orientation="h",
            yanchor="bottom",
            y=-0.2,
            xanchor="center",
            x=0.5
        )
    )
    
    return fig

def create_skills_bar_chart(matched_skills, missing_skills):
    skills_data = []
    for skill in matched_skills[:10]:
        skills_data.append({'Skill': skill, 'Status': 'Matched', 'Value': 1})
    for skill in missing_skills[:10]:
        skills_data.append({'Skill': skill, 'Status': 'Missing', 'Value': 1})
    
    if not skills_data:
        return None
    
    df = pd.DataFrame(skills_data)
    
    fig = px.bar(
        df,
        y='Skill',
        x='Value',
        color='Status',
        orientation='h',
        title='Top Skills Comparison',
        color_discrete_map={'Matched': '#667eea', 'Missing': '#f5576c'},
        height=400
    )
    
    fig.update_layout(
        paper_bgcolor='rgba(0,0,0,0)',
        plot_bgcolor='rgba(0,0,0,0)',
        font={'family': "Arial", 'color': '#333'},
        showlegend=True,
        xaxis_title="",
        yaxis_title="",
        xaxis={'showticklabels': False}
    )
    
    return fig

# Export functions
def export_to_pdf(results):
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Arial", 'B', 16)
    pdf.cell(0, 10, 'Resume Screening Analysis Report', 0, 1, 'C')
    pdf.ln(5)
    
    pdf.set_font("Arial", '', 12)
    pdf.cell(0, 10, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", 0, 1)
    pdf.ln(5)
    
    # Candidate Details
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, 'Candidate Information', 0, 1)
    pdf.set_font("Arial", '', 12)
    candidate = results['candidate_details']
    pdf.multi_cell(0, 10, f"Name: {candidate['name']}")
    pdf.multi_cell(0, 10, f"Email: {candidate['email']}")
    pdf.multi_cell(0, 10, f"Phone: {candidate['phone']}")
    pdf.multi_cell(0, 10, f"Experience: {candidate['experience']}")
    pdf.ln(5)
    
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, 'Overall Results', 0, 1)
    pdf.set_font("Arial", '', 12)
    pdf.multi_cell(0, 10, f"Match Percentage: {results['match_percentage']:.2f}%")
    pdf.multi_cell(0, 10, f"Recommendation: {results['recommendation']}")
    pdf.ln(5)
    
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, 'Matched Skills', 0, 1)
    pdf.set_font("Arial", '', 12)
    pdf.multi_cell(0, 10, ', '.join(results['matched_skills']) if results['matched_skills'] else 'None')
    pdf.ln(5)
    
    pdf.set_font("Arial", 'B', 14)
    pdf.cell(0, 10, 'Missing Skills', 0, 1)
    pdf.set_font("Arial", '', 12)
    pdf.multi_cell(0, 10, ', '.join(results['missing_skills']) if results['missing_skills'] else 'None')
    
    return pdf.output(dest='S').encode('latin1')

def export_to_txt(results):
    candidate = results['candidate_details']
    content = f"""
RESUME SCREENING ANALYSIS REPORT
{'='*50}

Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

CANDIDATE INFORMATION
{'-'*50}
Name: {candidate['name']}
Email: {candidate['email']}
Phone: {candidate['phone']}
Experience: {candidate['experience']}
LinkedIn: {candidate['linkedin']}
GitHub: {candidate['github']}

OVERALL RESULTS
{'-'*50}
Match Percentage: {results['match_percentage']:.2f}%
Recommendation: {results['recommendation']}

MATCHED SKILLS ({len(results['matched_skills'])})
{'-'*50}
{', '.join(results['matched_skills']) if results['matched_skills'] else 'None'}

MISSING SKILLS ({len(results['missing_skills'])})
{'-'*50}
{', '.join(results['missing_skills']) if results['missing_skills'] else 'None'}

TOP KEYWORDS
{'-'*50}
{', '.join(results['keywords']) if results['keywords'] else 'None'}
"""
    return content

def export_to_csv(results):
    candidate = results['candidate_details']
    data = {
        'Metric': [
            'Candidate Name',
            'Email',
            'Phone',
            'Experience',
            'Match Percentage',
            'Recommendation',
            'Matched Skills Count',
            'Missing Skills Count',
            'Total Resume Skills',
            'LinkedIn',
            'GitHub'
        ],
        'Value': [
            candidate['name'],
            candidate['email'],
            candidate['phone'],
            candidate['experience'],
            f"{results['match_percentage']:.2f}%",
            results['recommendation'],
            len(results['matched_skills']),
            len(results['missing_skills']),
            len(results['resume_skills']),
            candidate['linkedin'],
            candidate['github']
        ]
    }
    df = pd.DataFrame(data)
    return df.to_csv(index=False)

# Main UI
def main():
    # Header
    st.markdown('<h1 class="header-title">🎯 AI Resume Screener</h1>', unsafe_allow_html=True)
    st.markdown('<p class="subtitle"><strong> Powered by Advanced BERT Technology </strong></p>', unsafe_allow_html=True)
    
    # Sidebar
    with st.sidebar:
        st.markdown("### ⚙️ Settings")
        st.markdown("---")
        st.markdown("### 📊 About")
        st.info(
            "This AI-powered resume screener uses advanced NLP technology (BERT) "
            "to analyze resume-job description compatibility and provide detailed insights."
        )
        st.markdown("### 🎨 Features")
        st.success(
            "✅ Multi-format support (PDF, DOCX, TXT)\n\n"
            "✅ Semantic similarity analysis\n\n"
            "✅ Skill matching & extraction\n\n"
            "✅ Interactive visualizations\n\n"
            "✅ Export results (PDF, TXT, CSV)"
        )
    
    # Main content
    # st.markdown('<div class="upload-box">', unsafe_allow_html=True)
    st.markdown("<br>", unsafe_allow_html=True)
    
    col1, col2 = st.columns(2)
    
    with col1:
        st.markdown("### 📄 Upload Resume")
       
        resume_file = st.file_uploader(
            "Choose a resume file",
            type=['pdf', 'docx', 'txt'],
            key='resume',
            help="Upload candidate's resume in PDF, DOCX, or TXT format"
        )
    
    with col2:
        st.markdown("### 💼 Job Description")
        job_input_method = st.radio(
            "Input method:",
            ["Upload File", "Paste Text"],
            horizontal=True
        )
        
        if job_input_method == "Upload File":
            job_file = st.file_uploader(
                "Choose a job description file",
                type=['pdf', 'docx', 'txt'],
                key='job',
                help="Upload job description in PDF, DOCX, or TXT format"
            )
            job_description = extract_text(job_file) if job_file else ""
        else:
            job_description = st.text_area(
                "Paste job description here:",
                height=200,
                placeholder="Enter the job description..."
            )
    
    st.markdown('</div>', unsafe_allow_html=True)
    
    # Analyze button
    st.markdown("<br>", unsafe_allow_html=True)
    col_btn1, col_btn2, col_btn3 = st.columns([1, 1, 1])
    with col_btn2:
        
        analyze_button = st.button(" ## Analyze Resume", use_container_width=True)
    
    if analyze_button:
        if resume_file and job_description:
            with st.spinner('🔄 Analyzing resume with AI...'):
                resume_text = extract_text(resume_file)
                results = analyze_resume(resume_text, job_description)
                st.session_state.results = results
                st.session_state.analysis_done = True
        else:
            st.error("⚠️ Please upload a resume and provide a job description!")
    
    # Display results
    if st.session_state.analysis_done and st.session_state.results:
        results = st.session_state.results
        
        st.markdown("<br><br>", unsafe_allow_html=True)
        # st.markdown("## 📊 Analysis Results")
        st.markdown("---")
        
        # Create tabs
        tab1, tab2, tab3, tab4 = st.tabs(["📈 Analysis Results", "📊 Dashboard", "👤 Candidate Details", "💾 Export Reports"])
        
        # TAB 1: Analysis Results
        with tab1:
            # st.markdown('<div class="tab-content">', unsafe_allow_html=True)
            # st.markdown("---")
            # Metrics row
            metric_col1, metric_col2, metric_col3, metric_col4 = st.columns(4)
            
            with metric_col1:
                st.markdown(f"""
                    <div class="metric-card">
                        <h3 style="margin:0;">Match Score</h3>
                        <h1 style="margin:10px 0;">{results['match_percentage']:.1f}%</h1>
                    </div>
                """, unsafe_allow_html=True)
            
            with metric_col2:
                st.markdown(f"""
                    <div class="metric-card" style="background: linear-gradient(135deg, #4facfe 0%, #00f2fe 100%);">
                        <h3 style="margin:0;">Matched Skills</h3>
                        <h1 style="margin:10px 0;">{len(results['matched_skills'])}</h1>
                    </div>
                """, unsafe_allow_html=True)
            
            with metric_col3:
                st.markdown(f"""
                    <div class="metric-card" style="background: linear-gradient(135deg, #fa709a 0%, #fee140 100%);">
                        <h3 style="margin:0;">Missing Skills</h3>
                        <h1 style="margin:10px 0;">{len(results['missing_skills'])}</h1>
                    </div>
                """, unsafe_allow_html=True)
            
            with metric_col4:
                st.markdown(f"""
                    <div class="metric-card" style="background: linear-gradient(135deg, #30cfd0 0%, #330867 100%);">
                        <h3 style="margin:0;">Status</h3>
                        <h2 style="margin:10px 0;">{results['recommendation']}</h2>
                    </div>
                """, unsafe_allow_html=True)
            
            st.markdown("<br>", unsafe_allow_html=True)
            
            # Detailed results
            detail_col1, detail_col2 = st.columns(2)
            
            with detail_col1:
                # st.markdown('<div class="result-card">', unsafe_allow_html=True)
                st.markdown("### ✅ Matched Skills")
                
                if results['matched_skills']:
                    html = '<div style="display: flex; flex-wrap: wrap; gap: 8px;">'
                    for skill in results['matched_skills']:
                        # st.markdown(f'<span class="skill-badge">{skill}</span>', unsafe_allow_html=True)
                         html += f'<span style="background: linear-gradient(135deg, #4facfe 0%, #00f2fe 100%); color: white; padding: 5px 15px; border-radius: 20px; font-size: 14px;">{skill}</span>'
                    html += '</div>' 
                    st.markdown(html, unsafe_allow_html=True)      
                        # st.markdown(f'<span class="skill-badge">{skill}</span>', unsafe_allow_html=True)
                else:
                    st.info("No matched skills found")
                st.markdown('</div>', unsafe_allow_html=True)
            
            with detail_col2:
                # st.markdown('<div class="result-card">', unsafe_allow_html=True)
                st.markdown("### ❌ Missing Skills")
                if results['missing_skills']:
                    html = '<div style="display: flex; flex-wrap: wrap; gap: 8px;">'
                    for skill in results['missing_skills']:
                        html += f'<span style="background: linear-gradient(135deg, #fa709a 0%, #fee140 100%); color: white; padding: 5px 15px; border-radius: 20px; font-size: 14px;">{skill}</span>'
                        # st.markdown(f'<span class="skill-badge" style="background: linear-gradient(135deg, #f5576c 0%, #f093fb 100%);">{skill}</span>', unsafe_allow_html=True)
                        html += '</div>'
                        st.markdown(html, unsafe_allow_html=True)
                else:
                    st.success("All required skills are present!")
                st.markdown('</div>', unsafe_allow_html=True)
            
            # Top Keywords
            # st.markdown('<div class="result-card">', unsafe_allow_html=True)
            st.markdown("---", unsafe_allow_html=True)
            st.markdown("### 🔑 Top Keywords")
            if results['keywords']:
                keyword_text = " • ".join(results['keywords'])
                st.markdown(f"**{keyword_text}**")
            else:
                st.info("No significant keywords identified")
            st.markdown('</div>', unsafe_allow_html=True)
            
            # Recommendation Summary
            # st.markdown('<div class="result-card">', unsafe_allow_html=True)
            st.markdown("---", unsafe_allow_html=True)
            st.markdown("###")
            st.markdown(f"""
                ### Based on the AI analysis, this candidate is **{results['recommendation']}** for the position.
                
                - **Overall Match**: {results['match_percentage']:.1f}%
                - **Skills Coverage**: {len(results['matched_skills'])} out of {len(results['job_skills'])} required skills
                - **Skill Gap**: {len(results['missing_skills'])} skills need development
            """)
            st.markdown('</div>', unsafe_allow_html=True)
            
            st.markdown('</div>', unsafe_allow_html=True)
        
        # TAB 2: Dashboard
        with tab2:
            # st.markdown('<div class="tab-content">', unsafe_allow_html=True)
            st.markdown("### 📊 Visual Analytics Dashboard")
            
            # Charts row 1
            chart_col1, chart_col2 = st.columns(2)
            
            with chart_col1:
                st.plotly_chart(
                    create_gauge_chart(results['match_percentage'], "Overall Match Score"),
                    use_container_width=True
                )
            
            with chart_col2:
                st.plotly_chart(
                    create_skill_comparison_chart(results['matched_skills'], results['missing_skills']),
                    use_container_width=True
                )
            
            # Skills bar chart
            skills_chart = create_skills_bar_chart(results['matched_skills'], results['missing_skills'])
            if skills_chart:
                st.plotly_chart(skills_chart, use_container_width=True)
            
            # Additional metrics
            st.markdown("### 📈 Key Performance Indicators")
            kpi_col1, kpi_col2, kpi_col3 = st.columns(3)
            
            with kpi_col1:
                skill_match_rate = (len(results['matched_skills']) / len(results['job_skills']) * 100) if results['job_skills'] else 0
                st.metric("Skill Match Rate", f"{skill_match_rate:.1f}%", delta=f"{skill_match_rate - 60:.1f}%" if skill_match_rate > 60 else f"{skill_match_rate - 60:.1f}%")
            
            with kpi_col2:
                st.metric("Total Skills Found", len(results['resume_skills']), delta=f"+{len(results['matched_skills'])} matched")
            
            with kpi_col3:
                st.metric("Skill Gap", len(results['missing_skills']), delta=f"{len(results['missing_skills'])} to learn", delta_color="inverse")
            
            st.markdown('</div>', unsafe_allow_html=True)
        
        # TAB 3: Candidate Details
        with tab3:
            # st.markdown('<div class="tab-content">', unsafe_allow_html=True)
            st.markdown("### 👤 Candidate Profile")
            st.markdown("<br>", unsafe_allow_html=True)
            
            candidate = results['candidate_details']
            
            # Personal Information
            st.markdown("#### 📋 Personal Information")
            info_col1, info_col2 = st.columns(2)
            
            with info_col1:
                st.markdown(f"""
                    <div class="result-card">
                        <p><strong>📛 Name:</strong> {candidate['name']}</p>
                        <p><strong>📧 Email:</strong> {candidate['email']}</p>
                        <p><strong>📱 Phone:</strong> {candidate['phone']}</p>
                    </div>
                """, unsafe_allow_html=True)
            
            with info_col2:
                st.markdown(f"""
                    <div class="result-card">
                        <p><strong>💼 Experience:</strong> {candidate['experience']}</p>
                        <p><strong>📄 Resume Length:</strong> {candidate['word_count']} words</p>
                        <p><strong>🔗 LinkedIn:</strong> {candidate['linkedin']}</p>
                    </div>
                """, unsafe_allow_html=True)
            
            # Education
            st.markdown("#### 🎓 Education")
            st.markdown('<br>', unsafe_allow_html=True)
            # st.markdown('<div class="result-card">', unsafe_allow_html=True)
            if candidate['education'] != ['Not found']:
                for edu in candidate['education']:
                    st.markdown(f"• {edu}")
            else:
                st.info("No education information found")
            st.markdown('</div>', unsafe_allow_html=True)
            
            # Skills Profile
            # st.markdown("#### 🛠️ Skills Profile")
            st.markdown("""
                        <style>
                        .result-card {
                            background: radial-gradient(circle,rgba(238, 174, 202, 1) 0%, rgba(148, 187, 233, 1) 100%);
                            padding: 5px;
                            border-radius: 10px;
                            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.1);
                        }
                            
                        </style>
                        <div class="result-card">
                            <h4 style="margin-top:0;">🛠️ Skills Profile</h4>
                        </div>
                        """, unsafe_allow_html=True)
            skill_col1, skill_col2 = st.columns(2)
            
            with skill_col1:
                # st.markdown('<div class="result-card">', unsafe_allow_html=True)
                st.markdown("**✅ Possessed Skills**")
                if results['resume_skills']:
                    html = '<div style="display: flex; flex-wrap: wrap; gap:8px;">'
                
                    for skill in results['resume_skills'][:15]:
                    # st.markdown(f'<span class="skill-badge" style="background: linear-gradient(135deg, #4facfe 0%, #00f2fe 100%);">{skill}</span>', unsafe_allow_html=True)
                        html += f'<span style="background: linear-gradient(135deg, #4facfe 0%, #00f2fe 100%); color: white; padding: 5px 15px; border-radius: 20px; font-size: 14px;">{skill}</span>'
                    html += '</div>'
                    st.markdown(html, unsafe_allow_html=True)
                else:
                    st.info("No skills identified in resume")
                st.markdown('</div>', unsafe_allow_html=True)
            
            with skill_col2:
                # st.markdown('<div class="result-card">', unsafe_allow_html=True)
                st.markdown("**🎯 Job Required Skills**")
                if results['job_skills']:
                    html = '<div style="display: flex; flex-wrap: wrap; gap:8px;">'
                    for skill in results['job_skills'][:15]:
                        color = "linear-gradient(135deg, #00C853 0%, #64dd17 100%)" if skill in results['matched_skills'] else "linear-gradient(135deg, #fa709a 0%, #fee140 100%)"
                        html += f'<span style="background: {color}; color: white; padding: 5px 15px; border-radius: 20px; font-size: 14px;">{skill}</span>'
                    html += '</div>'
                    st.markdown(html, unsafe_allow_html=True)
                else:
                    st.info("No job skills identified")
                st.markdown('</div>', unsafe_allow_html=True)
            
            # Social Links
            # st.markdown("#### 🌐 Online Presence")
            # st.markdown('<div class="result-card">', unsafe_allow_html=True)
            st.markdown("""
                        <style>
                        .result-card {
                            background: linear-gradient(90deg,rgba(2, 0, 36, 1) 0%, rgba(9, 9, 121, 1) 84%, rgba(0, 212, 255, 1) 100%);
                            padding: 5px;
                            border-radius: 10px;
                            box-shadow: 0 4px 15px rgba(0, 0, 0, 0.1);
                        }
                            
                        </style>
                        <div class="result-card">
                            <h4 style="margin-top:0;">🌐 Online Presence</h4>
                        </div>
                        """, unsafe_allow_html=True)
            link_col1, link_col2 = st.columns(2)
            with link_col1:
                if candidate['linkedin'] != 'Not found':
                    st.markdown(f"🔗 **LinkedIn:** [{candidate['linkedin']}](https://{candidate['linkedin']})")
                else:
                    st.markdown("🔗 **LinkedIn:** Not provided")
            with link_col2:
                if candidate['github'] != 'Not found':
                    st.markdown(f"💻 **GitHub:** [{candidate['github']}](https://{candidate['github']})")
                else:
                    st.markdown("💻 **GitHub:** Not provided")
            st.markdown('</div>', unsafe_allow_html=True)
            
            st.markdown('</div>', unsafe_allow_html=True)
        
        # TAB 4: Export Reports
        with tab4:
            # st.markdown('<div class="tab-content">', unsafe_allow_html=True)
            st.markdown("### 💾 Export Analysis Reports")
            st.markdown("Download the complete analysis in your preferred format:")
            
            st.markdown("<br>", unsafe_allow_html=True)
            
            # Export options with descriptions
            exp_col1, exp_col2, exp_col3 = st.columns(3)
            
            with exp_col1:
                # st.markdown('<div class="result-card" style="text-align: center;">', unsafe_allow_html=True)
                st.markdown("### 📄 PDF Report")
                st.markdown("Professional formatted report with all analysis details")
                pdf_data = export_to_pdf(results)
                st.download_button(
                    label="📥 Download PDF",
                    data=pdf_data,
                    file_name=f"resume_analysis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.pdf",
                    mime="application/pdf",
                    use_container_width=True
                )
                st.markdown('</div>', unsafe_allow_html=True)
            
            with exp_col2:
                # st.markdown('<div class="result-card" style="text-align: center;">', unsafe_allow_html=True)
                st.markdown("### 📝 Text Report")
                st.markdown("Simple text format for easy sharing and editing")
                txt_data = export_to_txt(results)
                st.download_button(
                    label="📥 Download TXT",
                    data=txt_data,
                    file_name=f"resume_analysis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.txt",
                    mime="text/plain",
                    use_container_width=True
                )
                st.markdown('</div>', unsafe_allow_html=True)
            
            with exp_col3:
                # st.markdown('<div class="result-card" style="text-align: center;">', unsafe_allow_html=True)
                st.markdown("### 📊 CSV Data")
                st.markdown("Spreadsheet format for data analysis and tracking")
                csv_data = export_to_csv(results)
                st.download_button(
                    label="📥 Download CSV",
                    data=csv_data,
                    file_name=f"resume_analysis_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv",
                    mime="text/csv",
                    use_container_width=True
                )
                st.markdown('</div>', unsafe_allow_html=True)
            
            st.markdown("---", unsafe_allow_html=True)
            
            # Export preview
            st.markdown("### 👁️ Export Preview")
            preview_format = st.selectbox("Select format to preview:", ["Text", "CSV Data"])
            
            if preview_format == "Text":
                # st.markdown('<div class="result-card">', unsafe_allow_html=True)
                st.code(export_to_txt(results), language="text")
                st.markdown('</div>', unsafe_allow_html=True)
            else:
                # st.markdown('<div class="result-card">', unsafe_allow_html=True)
                csv_df = pd.read_csv(io.StringIO(export_to_csv(results)))
                st.dataframe(csv_df, use_container_width=True)
                st.markdown('</div>', unsafe_allow_html=True)
            
            st.markdown('</div>', unsafe_allow_html=True)

if __name__ == "__main__":
    main()
